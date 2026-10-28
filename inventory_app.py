"""
Inventory Health dashboard.

Runs sql/inventory_classifier.sql against the P21 database, caches the result,
keeps a snapshot history for trends, and serves web/dashboard.html.

Usage:
  python inventory_app.py                    start the dashboard (http://127.0.0.1:8050)
  python inventory_app.py --refresh          run the query once, update cache + history, exit
  python inventory_app.py --export FILE.html write a self-contained HTML report from the cache
  python inventory_app.py --refresh --export FILE.html   both, in that order (good for Task Scheduler)
  python inventory_app.py --test-connection  check the database connection and exit
  python inventory_app.py --demo             fill the cache with made-up data to try the dashboard
"""
from __future__ import annotations

import argparse
import configparser
import csv
import datetime as dt
import decimal
import io
import json
import os
import random
import re
import sys
import threading
import time
import traceback
import webbrowser
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
SQL_FILE = ROOT / "sql" / "inventory_classifier.sql"
HTML_FILE = ROOT / "web" / "dashboard.html"
DATA_DIR = ROOT / "data"
CACHE_JSON = DATA_DIR / "latest.json"
CACHE_XLSX = DATA_DIR / "latest.xlsx"
HISTORY_CSV = DATA_DIR / "history.csv"
CONFIG_FILE = ROOT / "config.ini"
DEMO_MARKER = DATA_DIR / "DEMO_DATA"
VENDOR_TERMS = DATA_DIR / "vendor_terms.json"   # rebates and terms typed into the vendor scorecard
SAMPLE_DIR = ROOT / "data_demo"                  # made-up dataset for the hidden sample mode (demo_data.py)

# Per-dashboard settings ([app] and [locations] in the config file). Defaults = the Abatix dashboard.
APP = {"title": "Inventory Health", "tagline": "Purchasing analytics",
       "corp_buyers": ["ZEVELY", "GAMBRELL", "RITTER", "MCCARDELL", "LUND", "LEER"],
       "managers": ["GAMBRELL"],
       "feedback_email": "michael.gambrell@abatix.com"}
INSTANCE_LOCS = None   # list of {"location_id","location_name","loc_type"} when the config defines [locations]

# Result sets in the order the SQL file returns them.
RESULT_SETS = [
    ("A", "Item Rollup"),
    ("B", "Top N by Group"),
    ("C", "Top Offenders"),
    ("D", "By Product Group"),
    ("E", "By Supplier"),
    ("F", "By Location"),
    ("G", "Status Mix"),
    ("H", "Loc Status Stockable"),
    ("I", "DOH Item"),
    ("J", "DOH Rollup"),
    ("K", "Transfer Usage"),
    ("L", "Run Info"),
    ("M", "Open Stock by Buyer"),
]

# Query parameters: config key -> (SQL variable, min, max)
PARAMS = {
    "company_id": ("CompanyId", 0, 10**9),
    "dead_days": ("DeadDays", 1, 5000),
    "slow_days": ("SlowDays", 1, 5000),
    "aging_days": ("AgingDays", 1, 5000),
    "usage_months": ("UsageMonths", 1, 36),
    "top_n_per_group": ("TopN_PerGroup", 1, 1000),
    "top_n_overall": ("TopN_Overall", 1, 5000),
    "change_days": ("ChangeDays", 7, 365),   # add-on modules only (inventory change window)
    "require_max": ("RequireMax", 0, 1),     # add-on modules only: 1 = an item needs max > 0 to count as stocked, 0 = min > 0 or max > 0
}

HIST_COLS = ["snapshot_date", "location_id", "location_name", "loc_type",
             "inv_status", "item_count", "qty_remaining", "cost_on_hand"]


# --------------------------------------------------------------------------- config

def apply_instance(cfg):
    """Point the app at this dashboard's data folder, name and location list."""
    global DATA_DIR, CACHE_JSON, CACHE_XLSX, HISTORY_CSV, DEMO_MARKER, INSTANCE_LOCS, VENDOR_TERMS
    if cfg.has_section("app"):
        APP["title"] = cfg.get("app", "title", fallback=APP["title"]).strip() or APP["title"]
        APP["tagline"] = cfg.get("app", "tagline", fallback=APP["tagline"]).strip()
        dd = cfg.get("app", "data_dir", fallback="").strip()
        if dd:
            DATA_DIR = ROOT / dd
            CACHE_JSON, CACHE_XLSX = DATA_DIR / "latest.json", DATA_DIR / "latest.xlsx"
            HISTORY_CSV, DEMO_MARKER = DATA_DIR / "history.csv", DATA_DIR / "DEMO_DATA"
            VENDOR_TERMS = DATA_DIR / "vendor_terms.json"
    team = cfg.get("purchasing", "corp_buyers", fallback="").strip() if cfg.has_section("purchasing") else ""
    fb = cfg.get("purchasing", "feedback_email", fallback="").strip() if cfg.has_section("purchasing") else ""
    if fb and "@" in fb:
        APP["feedback_email"] = fb
    if team:
        APP["corp_buyers"] = [t.strip().upper() for t in team.split(",") if re.fullmatch(r"[A-Za-z .'-]+", t.strip())]
    if cfg.has_section("purchasing") and cfg.has_option("purchasing", "managers"):   # blank = nobody excluded
        APP["managers"] = [t.strip().upper() for t in cfg.get("purchasing", "managers").split(",") if re.fullmatch(r"[A-Za-z .'-]+", t.strip())]
    if cfg.has_section("locations"):
        locs = []
        for k, v in cfg.items("locations"):
            if not k.strip().isdigit():
                continue
            name, _, kind = v.partition("|")
            locs.append({"location_id": int(k), "location_name": name.strip(),
                         "loc_type": "CORP" if kind.strip().lower() == "corp" else "BRANCH"})
        INSTANCE_LOCS = locs or None


def load_config() -> configparser.ConfigParser:
    cfg = configparser.ConfigParser(inline_comment_prefixes=(";",))
    cfg.read(CONFIG_FILE, encoding="utf-8")
    for section in ("database", "query", "server"):
        if not cfg.has_section(section):
            cfg.add_section(section)
    return cfg


def default_params(cfg) -> dict:
    fallback = {"company_id": 1, "dead_days": 365, "slow_days": 180, "aging_days": 90,
                "usage_months": 6, "top_n_per_group": 5, "top_n_overall": 25, "change_days": 61, "require_max": 1}
    return {k: cfg.getint("query", k, fallback=v) for k, v in fallback.items()}


def clean_params(raw: dict, base: dict) -> dict:
    """Validate user-supplied parameters. Values are forced to int, so they are safe to put in SQL."""
    out = dict(base)
    for key, (_, lo, hi) in PARAMS.items():
        if key in raw and raw[key] not in (None, ""):
            try:
                val = int(float(raw[key]))
            except (TypeError, ValueError):
                raise ValueError(f"{key} must be a whole number")
            if not lo <= val <= hi:
                raise ValueError(f"{key} must be between {lo} and {hi}")
            out[key] = val
    if not out["aging_days"] < out["slow_days"] < out["dead_days"]:
        raise ValueError("Thresholds must satisfy aging < slow < dead days")
    return out


def connection_string(cfg) -> str:
    db = cfg["database"]
    parts = []
    if db.get("dsn", "").strip():
        parts.append(f"DSN={db['dsn'].strip()}")
    else:
        parts.append(f"DRIVER={{{db.get('driver', 'ODBC Driver 17 for SQL Server').strip()}}}")
        parts.append(f"SERVER={db.get('server', 'SQL').strip()}")
    if db.get("database", "").strip():
        parts.append(f"DATABASE={db['database'].strip()}")
    if db.getboolean("trusted_connection", fallback=True):
        parts.append("Trusted_Connection=yes")
    else:
        parts.append("UID=" + odbc_quote(db.get("username", "").strip()))
        parts.append("PWD=" + odbc_quote(get_password(cfg)))
    parts.append("APP=Inventory Health Dashboard")
    return ";".join(parts) + ";"


def odbc_quote(v: str) -> str:
    return "{" + v.replace("}", "}}") + "}" if any(c in v for c in ";{}= ") else v


# --- SQL login password: config.ini, env var INVDB_PASSWORD, or encrypted file (Windows DPAPI) ---
PASSWORD_FILE = DATA_DIR / "db_password.bin"
_password_cache: str | None = None


def _dpapi(data: bytes, protect: bool) -> bytes:
    import ctypes
    from ctypes import wintypes

    class BLOB(ctypes.Structure):
        _fields_ = [("cbData", wintypes.DWORD), ("pbData", ctypes.POINTER(ctypes.c_char))]

    buf = ctypes.create_string_buffer(data, len(data))
    inp, out = BLOB(len(data), ctypes.cast(buf, ctypes.POINTER(ctypes.c_char))), BLOB()
    crypt32 = ctypes.windll.crypt32
    fn = crypt32.CryptProtectData if protect else crypt32.CryptUnprotectData
    fn.argtypes = [ctypes.POINTER(BLOB), ctypes.c_void_p, ctypes.c_void_p, ctypes.c_void_p,
                   ctypes.c_void_p, wintypes.DWORD, ctypes.POINTER(BLOB)]
    if not fn(ctypes.byref(inp), None, None, None, None, 0x1, ctypes.byref(out)):  # 0x1 = no UI
        raise OSError("Windows could not " + ("encrypt" if protect else "decrypt") + " the saved password")
    try:
        return ctypes.string_at(out.pbData, out.cbData)
    finally:
        ctypes.windll.kernel32.LocalFree(out.pbData)


def save_password(pw: str):
    DATA_DIR.mkdir(exist_ok=True)
    PASSWORD_FILE.write_bytes(_dpapi(pw.encode("utf-8"), True))


def get_password(cfg) -> str:
    global _password_cache
    if _password_cache is not None:
        return _password_cache
    pw = cfg.get("database", "password", fallback="") or os.environ.get("INVDB_PASSWORD", "")
    if not pw and PASSWORD_FILE.exists():
        pw = _dpapi(PASSWORD_FILE.read_bytes(), False).decode("utf-8")
    if not pw and sys.stdin and sys.stdin.isatty():
        import getpass
        pw = getpass.getpass(f"SQL password for {cfg.get('database', 'username', fallback='').strip()}: ")
    if not pw:
        raise RuntimeError("No SQL password available. Run: python inventory_app.py --save-password")
    _password_cache = pw
    return pw


def source_label(cfg) -> str:
    db = cfg["database"]
    where = f"DSN {db['dsn'].strip()}" if db.get("dsn", "").strip() else db.get("server", "SQL").strip()
    return f"{where} / {db.get('database', '').strip() or 'default database'}"


# --------------------------------------------------------------------------- SQL

def read_sql() -> str:
    return SQL_FILE.read_text(encoding="utf-8")


def prepare_sql(params: dict) -> str:
    sql = read_sql()
    for key, (var, _, _) in PARAMS.items():
        if key in MODULE_ONLY:
            continue
        pattern = re.compile(rf"(DECLARE\s+@{var}\s+\w+\s*=\s*)[^;]+;", re.IGNORECASE)
        sql, n = pattern.subn(lambda m: f"{m.group(1)}{int(params[key])};", sql, count=1)
        if n != 1:
            raise RuntimeError(f"Could not find 'DECLARE @{var} ... = value;' in {SQL_FILE.name}")
    if INSTANCE_LOCS:   # this dashboard's own corp / branch list replaces the #locs block
        vals = ",\n".join(f" ({l['location_id']}, '{l['location_name'].replace(chr(39), chr(39) * 2)}', {1 if l['loc_type'] == 'CORP' else 0})" for l in INSTANCE_LOCS)
        sql, n = re.subn(r"(INSERT INTO #locs \(location_id, location_name, is_corp\) VALUES\s*).*?;", lambda m: m.group(1) + vals + ";", sql, count=1, flags=re.S)
        if n != 1:
            raise RuntimeError("Could not find the INSERT INTO #locs block in the SQL file")
    if APP.get("corp_buyers"):   # corporate purchasing team from config replaces the #purch_buyers list
        vals = ", ".join("('" + t.replace("'", "''") + "')" for t in APP["corp_buyers"])
        sql, n = re.subn(r"(INSERT INTO #purch_buyers \(name_fragment\) VALUES\s*).*?;", lambda m: m.group(1) + vals + ";", sql, count=1, flags=re.S)
        if n != 1:
            raise RuntimeError("Could not find the INSERT INTO #purch_buyers block in the SQL file")
    # GO is an SSMS batch separator, not T-SQL.
    return re.sub(r"^\s*GO\s*;?\s*$", "", sql, flags=re.IGNORECASE | re.MULTILINE)


MODULE_DIR = ROOT / "sql" / "modules"
MODULE_PARAMS = {"company_id": "CompanyId", "slow_days": "SlowDays", "dead_days": "DeadDays", "aging_days": "AgingDays",
                 "change_days": "ChangeDays", "require_max": "RequireMax"}
MODULE_ONLY = {"change_days", "require_max"}


def module_files() -> list[Path]:
    return sorted(MODULE_DIR.glob("*.sql")) if MODULE_DIR.exists() else []


def prepare_module_sql(path: Path, params: dict) -> str:
    """Like prepare_sql, but only rewrites the parameters a module declares."""
    sql = path.read_text(encoding="utf-8")
    for key, var in MODULE_PARAMS.items():
        pattern = re.compile(rf"(DECLARE\s+@{var}\s+\w+\s*=\s*)[^;]+;", re.IGNORECASE)
        sql = pattern.sub(lambda m: f"{m.group(1)}{int(params[key])};", sql, count=1)
    if INSTANCE_LOCS:   # swap the default branch / company location lists for this dashboard's
        default = re.findall(r"\(\s*(\d+)\s*,\s*'[^']*'\s*,\s*([01])\s*\)", read_sql())
        d_branch = {int(i) for i, c in default if c == "0"}; d_all = {int(i) for i, _ in default}
        branch = ",".join(str(l["location_id"]) for l in INSTANCE_LOCS if l["loc_type"] == "BRANCH") or "-1"
        every = ",".join(str(l["location_id"]) for l in INSTANCE_LOCS)
        def swap(m):
            ids = {int(x) for x in re.findall(r"\d+", m.group(1))}
            return f"IN ({branch})" if ids == d_branch else f"IN ({every})" if ids == d_all else m.group(0)
        sql = re.sub(r"IN \(([\d,\s]+)\)", swap, sql)
    return re.sub(r"^\s*GO\s*;?\s*$", "", sql, flags=re.IGNORECASE | re.MULTILINE)


def run_one_module(cfg, path: Path, params: dict) -> dict:
    import pyodbc
    t0 = time.time()
    conn = pyodbc.connect(connection_string(cfg), autocommit=True, timeout=30)
    try:
        conn.timeout = cfg.getint("database", "query_timeout_seconds", fallback=0)
        cur = conn.cursor()
        cur.execute(prepare_module_sql(path, params))
        last = None
        while True:
            if cur.description:
                last = {"columns": [d[0] for d in cur.description],
                        "rows": [[to_json_value(v) for v in r] for r in cur.fetchall()]}
            if not cur.nextset():
                break
    finally:
        conn.close()
    if last is None:
        raise RuntimeError("returned no result set")
    last["seconds"] = round(time.time() - t0, 1)
    return last


def run_modules(cfg, params: dict, log=print, workers: int = 4) -> dict:
    """Run every sql/modules/*.sql in parallel. A failing module is recorded, never fatal."""
    from concurrent.futures import ThreadPoolExecutor
    files = module_files()
    if not files:
        return {}
    log(f"Running {len(files)} add-on queries ...")
    out = {}
    def job(f):
        try:
            res = run_one_module(cfg, f, params)
            log(f"  {f.stem}: {len(res['rows']):,} rows in {res['seconds']:.0f}s")
            return f.stem, res
        except Exception as ex:
            log(f"  {f.stem}: FAILED {friendly_error(ex)[:300]}")
            return f.stem, {"error": friendly_error(ex)}
    with ThreadPoolExecutor(max_workers=workers) as ex:
        for name, res in ex.map(job, files):
            out[name] = res
    return out


def sql_locations() -> list[dict]:
    """Read the #locs list from the SQL file so the location list lives in one place."""
    if INSTANCE_LOCS:
        return INSTANCE_LOCS
    rows = re.findall(r"\(\s*(\d+)\s*,\s*'([^']*)'\s*,\s*([01])\s*\)", read_sql())
    return [{"location_id": int(i), "location_name": n, "loc_type": "CORP" if c == "1" else "BRANCH"}
            for i, n, c in rows]


def to_json_value(v):
    if v is None or isinstance(v, (bool, int, float)):
        return v
    if isinstance(v, decimal.Decimal):
        f = float(v)
        return int(f) if f.is_integer() and abs(f) < 2**53 else f
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10]
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    return str(v).strip()


def run_query(cfg, params: dict, log=print) -> dict:
    import pyodbc  # imported here so --demo and --export work without it

    sql = prepare_sql(params)
    log(f"Connecting to {source_label(cfg)} ...")
    conn = pyodbc.connect(connection_string(cfg), autocommit=True, timeout=30)
    try:
        conn.timeout = cfg.getint("database", "query_timeout_seconds", fallback=0)
        cur = conn.cursor()
        log("Running classifier query ...")
        cur.execute(sql)
        sets = []
        while True:
            if cur.description:
                cols = [d[0] for d in cur.description]
                rows = [[to_json_value(v) for v in r] for r in cur.fetchall()]
                sets.append({"columns": cols, "rows": rows})
                log(f"  result set {len(sets)}: {len(rows):,} rows")
            if not cur.nextset():
                break
    finally:
        conn.close()
    if len(sets) != len(RESULT_SETS):
        raise RuntimeError(f"Expected {len(RESULT_SETS)} result sets from the query but got {len(sets)}. "
                           "If you edited the SQL, keep the outputs in the documented order.")
    return {code: s for (code, _), s in zip(RESULT_SETS, sets)}


# --------------------------------------------------------------------------- outputs

def build_payload(sets: dict, params: dict, source: str, seconds: float) -> dict:
    info_set = sets["L"]
    run_info = dict(zip(info_set["columns"], info_set["rows"][0])) if info_set["rows"] else {}
    return {
        "app": dict(APP),
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "source": source,
        "query_seconds": round(seconds, 1),
        "params": params,
        "run_info": run_info,
        "locations": sql_locations(),
        "items": sets["A"],
        "top_layers": sets["C"],
        "usage": sets["K"],
        "buyers": sets.get("M"),
        "modules": sets.get("_modules", {}),
    }


def write_workbook(sets: dict, path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook(write_only=True)
    for code, title in RESULT_SETS:
        if code not in sets:
            continue
        ws = wb.create_sheet(f"{code} {title}"[:31])
        ws.freeze_panes = "A2"
        header = []
        for c in sets[code]["columns"]:
            from openpyxl.cell import WriteOnlyCell
            cell = WriteOnlyCell(ws, value=c)
            cell.font = Font(bold=True)
            header.append(cell)
        ws.append(header)
        for row in sets[code]["rows"]:
            ws.append(row)
    for name, res in (sets.get("_modules") or {}).items():
        if "rows" not in res:
            continue
        ws = wb.create_sheet(("+ " + name.replace("_", " ").title())[:31])
        ws.freeze_panes = "A2"
        ws.append(res["columns"])
        for row in res["rows"]:
            ws.append(row)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def snapshot_rows(payload: dict) -> list[dict]:
    items = payload["items"]
    idx = {c: i for i, c in enumerate(items["columns"])}
    date = (payload.get("run_info") or {}).get("run_date") or dt.date.today().isoformat()
    agg: dict[tuple, dict] = {}
    for r in items["rows"]:
        key = (str(r[idx["location_id"]]), r[idx["inv_status"]])
        a = agg.setdefault(key, {
            "snapshot_date": date, "location_id": key[0], "location_name": r[idx["location_name"]],
            "loc_type": r[idx["loc_type"]], "inv_status": key[1],
            "item_count": 0, "qty_remaining": 0.0, "cost_on_hand": 0.0})
        a["item_count"] += 1
        a["qty_remaining"] += float(r[idx["qty_remaining"]] or 0)
        a["cost_on_hand"] += float(r[idx["cost_on_hand"]] or 0)
    for a in agg.values():
        a["qty_remaining"] = round(a["qty_remaining"], 4)
        a["cost_on_hand"] = round(a["cost_on_hand"], 2)
    return list(agg.values())


# ---- vendor terms (rebates, programs, notes) typed into the vendor scorecard ----
TERM_FIELDS = {   # field: (kind, limit)
    "rebate_pct": ("num", 100), "rebate_basis": ("text", 40), "rebate_threshold": ("num", 1e10), "rebate_period": ("text", 40),
    "growth_rebate_pct": ("num", 100), "growth_target_pct": ("num", 1000), "coop_pct": ("num", 100), "payment_terms": ("text", 60),
    "freight_terms": ("text", 80), "min_order": ("text", 80), "program_end": ("text", 20), "contact": ("text", 120), "notes": ("text", 2000),
}


def read_vendor_terms(path: Path | None = None) -> dict:
    try:
        data = json.loads((path or VENDOR_TERMS).read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def save_vendor_terms(supplier_id: str, raw: dict, user: str = "", path: Path | None = None) -> dict:
    """Validate and store one supplier's terms. Empty values are dropped; an empty set removes the supplier."""
    sid = str(supplier_id).strip()
    if not sid or len(sid) > 40:
        raise ValueError("Missing supplier id")
    clean = {}
    for key, (kind, limit) in TERM_FIELDS.items():
        v = raw.get(key)
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if kind == "num":
            v = float(v)
            if not (-limit <= v <= limit):
                raise ValueError(f"{key} is out of range")
            clean[key] = round(v, 4)
        else:
            clean[key] = str(v).strip()[: int(limit)]
    path = path or VENDOR_TERMS
    terms = read_vendor_terms(path)
    if clean:
        clean["supplier_name"] = str(raw.get("supplier_name", ""))[:120]
        clean["updated"] = dt.datetime.now().isoformat(timespec="seconds")
        clean["updated_by"] = user[:60]
        terms[sid] = clean
    else:
        terms.pop(sid, None)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(terms, indent=1), encoding="utf-8")
    tmp.replace(path)
    return terms


def read_history(path: Path | None = None) -> list[dict]:
    path = path or HISTORY_CSV
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [row for row in csv.DictReader(f) if row.get("snapshot_date")]


def update_history(payload: dict) -> int:
    """Replace any snapshot for the same date with the new one. Returns snapshot count."""
    new = snapshot_rows(payload)
    date = new[0]["snapshot_date"] if new else None
    rows = [r for r in read_history() if r["snapshot_date"] != date] + new
    rows.sort(key=lambda r: (r["snapshot_date"], str(r["location_id"]), r["inv_status"]))
    tmp = HISTORY_CSV.with_suffix(".tmp")
    with tmp.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HIST_COLS, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    tmp.replace(HISTORY_CSV)
    return len({r["snapshot_date"] for r in rows})


def save_workbook(sets: dict, log=print):
    try:
        write_workbook(sets, CACHE_XLSX)
        log(f"Saved {CACHE_XLSX.relative_to(ROOT)}")
    except Exception as ex:  # the dashboard still works without the workbook
        log(f"Could not write Excel workbook: {ex}")


def save_outputs(sets: dict, payload: dict, log=print, history=True, workbook=True):
    DATA_DIR.mkdir(exist_ok=True)
    tmp = CACHE_JSON.with_suffix(".tmp")
    tmp.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
    tmp.replace(CACHE_JSON)
    log(f"Saved {CACHE_JSON.relative_to(ROOT)}")
    if history:
        if DEMO_MARKER.exists():  # don't mix made-up snapshots into real history
            HISTORY_CSV.unlink(missing_ok=True)
            DEMO_MARKER.unlink()
            log("Cleared demo history")
        n = update_history(payload)
        log(f"History updated ({n} snapshot dates) in {HISTORY_CSV.relative_to(ROOT)}")
    if workbook:
        save_workbook(sets, log)


def load_cache() -> dict | None:
    if CACHE_JSON.exists():
        return json.loads(CACHE_JSON.read_text(encoding="utf-8"))
    return None


def render_static_html(payload: dict, files: dict | None = None) -> str:
    html = HTML_FILE.read_text(encoding="utf-8")
    files = files or {}
    bundle = {"payload": payload, "history": read_history(files.get("history")), "vendor_terms": read_vendor_terms(files.get("terms")), "static": True}
    data = json.dumps(bundle, separators=(",", ":")).replace("</", "<\\/")
    marker = '<script id="embedded-data" type="application/json">null</script>'
    if marker not in html:
        raise RuntimeError("dashboard.html is missing the embedded-data placeholder")
    return html.replace(marker, f'<script id="embedded-data" type="application/json">{data}</script>')


# --------------------------------------------------------------------------- demo data

def demo_sets(params: dict, seed: int = 7) -> dict:
    rnd = random.Random(seed)
    today = dt.date.today()
    locs = sql_locations() + [{"location_id": 90010, "location_name": "Consignment", "loc_type": "CONSIGNMENT"},
                              {"location_id": 90020, "location_name": "Consignment", "loc_type": "CONSIGNMENT"}]
    groups = ["Respirators", "Gloves", "Coveralls", "Absorbents", "Containment", "Tape & Adhesives",
              "Air Scrubbers", "Dehumidifiers", "Moisture Meters", "Hand Tools", "Sprayers",
              "Fire Restoration", "Mold Remediation", "Eye Protection", "Fasteners"]
    suppliers = [(1001, "3M Company"), (1002, "Kimberly-Clark"), (1003, "Ansell"), (1004, "DuPont"),
                 (1005, "Dri-Eaz"), (1006, "Abatement Technologies"), (1007, "Phoenix Restoration"),
                 (1008, "Honeywell Safety"), (1009, "Tramex"), (1010, "Shurtape"), (1011, "Chapin"),
                 (1012, "Benefect")]
    items = []
    for n in range(650):
        g = rnd.choice(groups)
        sid, sname = rnd.choice(suppliers)
        items.append((f"{g[:3].upper()}-{10000 + n}", f"{g[:-1] if g.endswith('s') else g} item {n} "
                      f"{rnd.choice(['XL', 'L', 'M', '12pk', 'case/100', 'each', '5 gal', 'roll'])}",
                      g, sid, sname, rnd.choice("AABBBCCCC"), round(rnd.lognormvariate(2.6, 1.1), 2)))
    A, K, M, layers = [], [], [], []
    purch_team = ["Jordan Avery", "Casey Brooks", "Morgan Ellis"]   # made-up names
    first = ["Sam", "Alex", "Taylor", "Jamie", "Drew", "Quinn", "Reese", "Skyler", "Parker", "Rowan", "Emerson",
             "Hayden", "Kendall", "Logan"]
    last = ["Ortiz", "Nguyen", "Patel", "Hughes", "Foster", "Ramos", "Bennett", "Diaz", "Coleman", "Price",
            "Sullivan", "Myers", "Hayes", "Ward"]
    branch_buyer = {l["location_id"]: f"{rnd.choice(first)} {rnd.choice(last)}" for l in locs}
    usage_map = {}
    for loc in locs:
        stocked = rnd.sample(items, rnd.randint(120, 260) if loc["loc_type"] == "BRANCH" else rnd.randint(30, 90))
        for (item_id, desc, grp, sid, sname, abc, unit) in stocked:
            r = rnd.random()
            dss = None if r < 0.08 else int(rnd.choice([rnd.uniform(0, 90)] * 5 + [rnd.uniform(90, 180)] * 2 +
                                                       [rnd.uniform(180, 365)] * 2 + [rnd.uniform(365, 1400)] * 2))
            qty = round(rnd.lognormvariate(3, 1.2))
            qty = max(qty, 1)
            cost_unit = unit * rnd.uniform(0.85, 1.1)
            cost = round(qty * cost_unit, 2)
            if loc["loc_type"] == "CONSIGNMENT":
                status, rank = "CONSIGNMENT", -1
            elif loc["loc_type"] == "CORP":
                status, rank = "CORP", 0
            elif dss is None or dss >= params["dead_days"]:
                status, rank = "DEAD", 4
            elif dss >= params["slow_days"]:
                status, rank = "SLOW", 3
            elif dss >= params["aging_days"]:
                status, rank = "AGING", 2
            else:
                status, rank = "ACTIVE", 1
            monthly = 0.0 if status in ("DEAD",) else max(0.0, rnd.lognormvariate(1.5, 1.2) * (0.3 if status == "SLOW" else 1))
            total_usage = round(monthly * params["usage_months"])
            work_days = int(params["usage_months"] * 30.4 * 5 / 7)
            doh = round(qty / (total_usage / work_days), 1) if total_usage else None
            oldest_days = (dss or rnd.randint(400, 1500)) + rnd.randint(0, 400)
            older = rnd.randint(1, 3) if oldest_days >= params["dead_days"] else 0
            older_qty = round(qty * rnd.uniform(0.3, 1)) if older else 0
            sale = (today - dt.timedelta(days=dss)).isoformat() if dss is not None else None
            recv_days = max(0, min(oldest_days, (dss or 500) + rnd.randint(-30, 200)))
            A.append([loc["location_id"], loc["location_name"], loc["loc_type"], item_id, desc, grp, abc, sid, sname,
                      rnd.choice("YYYYN"), qty, cost, round(unit, 4),
                      (today - dt.timedelta(days=recv_days)).isoformat(), sale,
                      (today - dt.timedelta(days=recv_days + 3)).isoformat(),
                      recv_days, dss, recv_days + 3, (today - dt.timedelta(days=oldest_days)).isoformat(), oldest_days,
                      older, older_qty, round(older_qty * cost_unit, 2), rank, status,
                      total_usage, round(total_usage / params["usage_months"], 4), doh])
            # buyer attribution: split the stock across 1-2 requesters
            if loc["loc_type"] == "CONSIGNMENT":
                picks = [(rnd.choice(["(Non-PO: TR)", "(Non-PO: IA)"]), "Non-PO/Trace", "Trace Needed")]
            else:
                picks = []
                for _ in range(rnd.choice([1, 1, 1, 2])):
                    r2 = rnd.random()
                    if r2 < 0.45:
                        picks.append((rnd.choice(purch_team), "Purchasing Department", "PO Origin"))
                    elif r2 < 0.9:
                        picks.append((branch_buyer[loc["location_id"]], "Branch", "PO Origin"))
                    else:
                        picks.append(("(Non-PO: TR)", "Non-PO/Trace", "Trace Needed"))
                picks = list(dict.fromkeys(picks))
            share = [rnd.uniform(0.3, 1) for _ in picks]
            tot_share = sum(share)
            last_po = None
            for (buyer, dept, origin), sh in zip(picks, share):
                po = rnd.randint(4_000_000, 4_999_999) if origin == "PO Origin" else None
                last_po = last_po or po
                M.append([loc["location_id"], item_id, buyer, dept, origin, rnd.randint(1, 3), 1 if po else 0, po,
                          round(qty * sh / tot_share, 4), round(cost * sh / tot_share, 2),
                          (today - dt.timedelta(days=oldest_days)).isoformat(),
                          (today - dt.timedelta(days=recv_days)).isoformat()])
            po_pick = next((p for p in picks if p[2] == "PO Origin"), None)
            A[-1] += [max(0, round(qty * rnd.uniform(0.4, 1.3))) if loc["loc_type"] != "CONSIGNMENT" else None,
                      po_pick[0] if po_pick else None, po_pick[1] if po_pick else None, last_po]
            if total_usage:
                usage_map[(item_id, loc["location_id"])] = total_usage
            layers.append([sid, sname, f"PO{rnd.randint(100000, 999999)}", (today - dt.timedelta(days=oldest_days)).isoformat(),
                           oldest_days, round(cost_unit, 4), round(unit, 4), qty, qty, cost, loc["location_id"],
                           loc["location_name"], loc["loc_type"], item_id, desc, abc, sale, None, grp, "Y",
                           status, rank])
    # Usage at branches that hold no stock (stock-outs) so transfers have somewhere to go.
    branches = [l for l in locs if l["loc_type"] == "BRANCH"]
    dead_items = {row[3] for row in A if row[25] in ("DEAD", "SLOW")}
    for item_id in dead_items:
        for loc in rnd.sample(branches, rnd.randint(0, 3)):
            usage_map.setdefault((item_id, loc["location_id"]), rnd.randint(4, 120))
    for (item_id, loc_id), u in usage_map.items():
        if item_id in dead_items:
            K.append([item_id, loc_id, u])
    cols_A = ["location_id", "location_name", "loc_type", "item_id", "description", "product_group", "abc_class",
              "supplier_id", "supplier_name", "stockable", "qty_remaining", "cost_on_hand", "replacement_cost",
              "most_recent_receipt_date", "most_recent_sale_date", "most_recent_purchase_date",
              "days_since_recent_receipt", "days_since_recent_sale", "days_since_recent_purchase",
              "oldest_open_layer_date", "max_days_since_received_layer", "older_layer_count",
              "qty_in_older_layers", "cost_in_older_layers", "status_rank", "inv_status",
              "total_usage", "avg_monthly_usage", "days_on_hand",
              "inv_max", "last_buyer", "last_buyer_dept", "last_po_no"]
    cols_M = ["location_id", "item_id", "buyer", "requester_dept", "origin_source", "layer_count", "po_count",
              "latest_po_no", "qty_remaining", "cost_on_hand", "first_received", "last_received"]
    cols_C = ["supplier_id", "supplier_name", "fifo_layer_no", "date_received", "days_since_received_layer",
              "fifo_layer_cost", "replacement_cost", "qty_received", "qty_remaining", "cost_on_hand", "location_id",
              "location_name", "loc_type", "item_id", "description", "abc_class", "last_sale_date",
              "last_purchase_date", "product_group", "stockable", "inv_status", "status_rank"]
    layers.sort(key=lambda r: -r[9])
    top = [r + [None, i + 1] for i, r in enumerate(layers[:params["top_n_overall"]])]
    ws = dt.date(today.year, today.month, 1)
    for _ in range(params["usage_months"]):
        ws = (ws - dt.timedelta(days=1)).replace(day=1)
    we = dt.date(today.year, today.month, 1) - dt.timedelta(days=1)
    info = [[today.isoformat(), ws.isoformat(), we.isoformat(), int(((we - ws).days + 1) * 5 / 7),
             params["usage_months"], params["company_id"], params["dead_days"], params["slow_days"], params["aging_days"]]]
    return {
        "A": {"columns": cols_A, "rows": A},
        "C": {"columns": cols_C + ["rn_pg_topN", "rn_overall"], "rows": top},
        "K": {"columns": ["item_id", "location_id", "total_usage"], "rows": K},
        "L": {"columns": ["run_date", "usage_window_start", "usage_window_end", "work_days", "usage_months",
                          "company_id", "dead_days", "slow_days", "aging_days"], "rows": info},
        "M": {"columns": cols_M, "rows": M},
    }


def demo_history(payload: dict):
    """Write a few fake past snapshots so the trend chart has something to show."""
    rnd = random.Random(3)
    base = snapshot_rows(payload)
    rows = []
    today = dt.date.fromisoformat(payload["run_info"]["run_date"])
    for weeks_back in range(10, 0, -1):
        d = (today - dt.timedelta(weeks=weeks_back)).isoformat()
        drift = 1 + weeks_back * 0.012
        for r in base:
            f = drift * rnd.uniform(0.97, 1.03) if r["inv_status"] in ("DEAD", "SLOW") else rnd.uniform(0.95, 1.05)
            rows.append({**r, "snapshot_date": d, "cost_on_hand": round(r["cost_on_hand"] * f, 2),
                         "qty_remaining": round(r["qty_remaining"] * f, 4)})
    rows += base
    DATA_DIR.mkdir(exist_ok=True)
    with HISTORY_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=HIST_COLS)
        w.writeheader()
        w.writerows(rows)


# --------------------------------------------------------------------------- refresh job

class RefreshJob:
    def __init__(self, cfg):
        self.cfg = cfg
        self.lock = threading.Lock()
        self.state = {"running": False, "started": None, "finished": None, "error": None,
                      "log": [], "params": None}

    def log(self, msg):
        stamp = dt.datetime.now().strftime("%H:%M:%S")
        line = f"[{stamp}] {msg}"
        print(line, flush=True)
        self.state["log"] = (self.state["log"] + [line])[-30:]

    def start(self, params: dict) -> bool:
        with self.lock:
            if self.state["running"]:
                return False
            self.state.update(running=True, started=time.time(), finished=None, error=None, log=[], params=params)
        threading.Thread(target=self._run, args=(params,), daemon=True).start()
        return True

    def _run(self, params):
        t0 = time.time()
        try:
            sets = run_query(self.cfg, params, log=self.log)
            sets["_modules"] = run_modules(self.cfg, params, log=self.log)
            payload = build_payload(sets, params, source_label(self.cfg), time.time() - t0)
            save_outputs(sets, payload, log=self.log, workbook=False)
            self.log(f"Done in {time.time() - t0:,.0f}s. Writing the Excel workbook in the background ...")
            threading.Thread(target=save_workbook, args=(sets, self.log), daemon=True).start()
        except Exception as ex:
            self.state["error"] = friendly_error(ex)
            self.log("FAILED: " + self.state["error"])
            traceback.print_exc()
        finally:
            self.state.update(running=False, finished=time.time())

    def status(self) -> dict:
        s = dict(self.state)
        s["elapsed"] = round((s["finished"] or time.time()) - s["started"], 1) if s["started"] else None
        return s


def friendly_error(ex: Exception) -> str:
    msg = str(ex)
    if "IM002" in msg:
        return "ODBC driver or DSN not found. Check driver/dsn in config.ini. " + msg
    if "08001" in msg:
        return "Could not reach the SQL Server. Check server in config.ini and your network/VPN. " + msg
    if "SSPI" in msg or "domain controller" in msg:
        return ("Windows sign-in to SQL Server failed because this PC can't reach a domain controller. "
                "Connect to the office network/VPN, or set trusted_connection = no with a SQL username/password "
                "in config.ini. " + msg)
    if "28000" in msg or "18456" in msg:
        return "Login failed. Check trusted_connection or username/password in config.ini. " + msg
    if "HYT00" in msg or "timeout" in msg.lower():
        return "The query timed out. Raise query_timeout_seconds in config.ini. " + msg
    return msg


# --------------------------------------------------------------------------- sample mode

def sample_files() -> dict:
    return {"json": SAMPLE_DIR / "latest.json", "xlsx": SAMPLE_DIR / "latest.xlsx",
            "history": SAMPLE_DIR / "history.csv", "terms": SAMPLE_DIR / "vendor_terms.json"}


_sample_lock = threading.Lock()


def ensure_sample(force: bool = False) -> dict:
    """Build the made-up dataset if it's missing or from an earlier day, so its dates stay current."""
    f = sample_files()
    with _sample_lock:
        stale = force or not f["json"].exists() or dt.date.fromtimestamp(f["json"].stat().st_mtime) != dt.date.today()
        if stale:
            import demo_data
            terms = read_vendor_terms(f["terms"]) if f["terms"].exists() else None
            demo_data.write(SAMPLE_DIR, workbook="background")
            if terms is not None:   # keep anything typed into the sample scorecard
                f["terms"].write_text(json.dumps(terms, indent=1), encoding="utf-8")
    return f


class SampleJob:
    """Pretend refresh for the sample mode: never touches the database, just rebuilds the made-up data."""
    STEPS = ["Connecting to {src}", "Running inventory classifier", "Classifier returned 5 result sets",
             "Running 19 add-on modules (4 at a time)", "Modules finished", "Saved latest.json", "History updated"]

    def __init__(self):
        self.started = None

    def start(self):
        if self.started and time.time() - self.started < 18:
            return False
        self.started = time.time()
        return True

    def status(self) -> dict:
        if not self.started:
            return {"running": False, "started": None, "finished": None, "error": None, "log": [], "elapsed": None}
        el = time.time() - self.started
        n = min(len(self.STEPS), 1 + int(el / 2.6))
        stamp = lambda i: dt.datetime.fromtimestamp(self.started + i * 2.6).strftime("%H:%M:%S")
        import demo_data
        log = [f"[{stamp(i)}] " + self.STEPS[i].format(src=demo_data.SOURCE) for i in range(n)]
        running = el < 18
        if not running:
            ensure_sample(force=dt.date.fromtimestamp(sample_files()["json"].stat().st_mtime) != dt.date.today())
        return {"running": running, "started": self.started, "finished": None if running else self.started + 18,
                "error": None, "log": log, "elapsed": round(min(el, 18), 1)}


# --------------------------------------------------------------------------- web server

class Handler(BaseHTTPRequestHandler):
    cfg = None
    job: RefreshJob = None
    sample_job = SampleJob()

    def _sample(self) -> bool:
        """Sample mode is chosen per request by the page (hidden switch), never stored on the server."""
        return self.headers.get("X-Data-Mode") == "sample" or "mode=sample" in (urlparse(self.path).query or "")

    def log_message(self, fmt, *args):  # keep the console quiet; refresh progress is logged separately
        pass

    def _send(self, body: bytes, ctype: str, status=HTTPStatus.OK, extra: dict | None = None):
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, status=HTTPStatus.OK):
        self._send(json.dumps(obj).encode("utf-8"), "application/json; charset=utf-8", status)

    def do_GET(self):
        path = urlparse(self.path).path
        try:
            if path in ("/", "/index.html"):
                self._send(HTML_FILE.read_bytes(), "text/html; charset=utf-8")
            elif path == "/api/data" and self._sample():
                self._send(ensure_sample()["json"].read_bytes(), "application/json; charset=utf-8")
            elif path == "/api/vendor_terms" and self._sample():
                self._json(read_vendor_terms(sample_files()["terms"]))
            elif path == "/api/history" and self._sample():
                self._json(read_history(ensure_sample()["history"]))
            elif path == "/api/status" and self._sample():
                self._json(self.sample_job.status())
            elif path == "/api/config" and self._sample():
                import demo_data
                self._json({"defaults": default_params(self.cfg), "source": demo_data.SOURCE, "app": APP})
            elif path == "/api/workbook" and self._sample():
                f = ensure_sample()
                if not f["xlsx"].exists():
                    return self._json({"error": "No workbook yet. Run a refresh first."}, HTTPStatus.NOT_FOUND)
                self._send(f["xlsx"].read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           extra={"Content-Disposition": f'attachment; filename="inventory_health_{dt.date.today().isoformat()}.xlsx"'})
            elif path == "/api/history.csv" and self._sample():
                self._send(ensure_sample()["history"].read_bytes(), "text/csv; charset=utf-8",
                           extra={"Content-Disposition": 'attachment; filename="inventory_history.csv"'})
            elif path == "/api/report.html" and self._sample():
                f = ensure_sample()
                self._send(render_static_html(json.loads(f["json"].read_text(encoding="utf-8")), f).encode("utf-8"), "text/html; charset=utf-8",
                           extra={"Content-Disposition": f'attachment; filename="inventory_health_{dt.date.today().isoformat()}.html"'})
            elif path == "/api/data":
                if CACHE_JSON.exists():
                    self._send(CACHE_JSON.read_bytes(), "application/json; charset=utf-8")
                else:
                    self._json({"error": "No data yet. Click Refresh to run the query."}, HTTPStatus.NOT_FOUND)
            elif path == "/api/vendor_terms":
                self._json(read_vendor_terms())
            elif path == "/api/history":
                self._json(read_history())
            elif path == "/api/status":
                self._json(self.job.status())
            elif path == "/api/config":
                self._json({"defaults": default_params(self.cfg), "source": source_label(self.cfg), "app": APP})
            elif path == "/api/workbook":
                if not CACHE_XLSX.exists():
                    return self._json({"error": "No workbook yet. Run a refresh first."}, HTTPStatus.NOT_FOUND)
                name = f"inventory_health_{dt.date.today().isoformat()}.xlsx"
                self._send(CACHE_XLSX.read_bytes(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           extra={"Content-Disposition": f'attachment; filename="{name}"'})
            elif path == "/api/history.csv":
                body = HISTORY_CSV.read_bytes() if HISTORY_CSV.exists() else (",".join(HIST_COLS) + "\r\n").encode()
                self._send(body, "text/csv; charset=utf-8",
                           extra={"Content-Disposition": 'attachment; filename="inventory_history.csv"'})
            elif path == "/api/report.html":
                payload = load_cache()
                if not payload:
                    return self._json({"error": "No data yet."}, HTTPStatus.NOT_FOUND)
                name = f"inventory_health_{dt.date.today().isoformat()}.html"
                self._send(render_static_html(payload).encode("utf-8"), "text/html; charset=utf-8",
                           extra={"Content-Disposition": f'attachment; filename="{name}"'})
            else:
                self._json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
        except Exception as ex:
            traceback.print_exc()
            self._json({"error": str(ex)}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def do_POST(self):
        path = urlparse(self.path).path
        if path == "/api/vendor_terms":
            try:
                length = int(self.headers.get("Content-Length") or 0)
                if length > 20000:
                    raise ValueError("Too much data")
                raw = json.loads(self.rfile.read(length) or b"{}")
                if self._sample():
                    terms = save_vendor_terms(raw.get("supplier_id", ""), raw.get("terms") or {}, "dwhitfield", sample_files()["terms"])
                else:
                    terms = save_vendor_terms(raw.get("supplier_id", ""), raw.get("terms") or {}, os.environ.get("USERNAME", ""))
            except (ValueError, TypeError, json.JSONDecodeError) as ex:
                return self._json({"error": str(ex)}, HTTPStatus.BAD_REQUEST)
            return self._json(terms)
        if path != "/api/refresh":
            return self._json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
        if self._sample():
            self.rfile.read(int(self.headers.get("Content-Length") or 0))
            if not self.sample_job.start():
                return self._json({"error": "A refresh is already running."}, HTTPStatus.CONFLICT)
            return self._json({"started": True})
        try:
            length = int(self.headers.get("Content-Length") or 0)
            raw = json.loads(self.rfile.read(length) or b"{}") if length else {}
            params = clean_params(raw, default_params(self.cfg))
        except (ValueError, json.JSONDecodeError) as ex:
            return self._json({"error": str(ex)}, HTTPStatus.BAD_REQUEST)
        if not self.job.start(params):
            return self._json({"error": "A refresh is already running."}, HTTPStatus.CONFLICT)
        self._json({"started": True, "params": params}, HTTPStatus.ACCEPTED)


def serve(cfg):
    host = cfg.get("server", "host", fallback="127.0.0.1")
    port = cfg.getint("server", "port", fallback=8050)
    Handler.cfg = cfg
    Handler.job = RefreshJob(cfg)
    httpd = ThreadingHTTPServer((host, port), Handler)
    url = f"http://{'127.0.0.1' if host in ('0.0.0.0', '') else host}:{port}/"
    print(f"{APP['title']} running at {url}")
    print(f"Database: {source_label(cfg)}")
    print("Close this window or press Ctrl+C to stop.")
    if cfg.getboolean("server", "open_browser", fallback=True):
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")


# --------------------------------------------------------------------------- CLI

def main(argv=None):
    ap = argparse.ArgumentParser(description="Inventory Health dashboard")
    ap.add_argument("--refresh", action="store_true", help="run the query once and update the cache")
    ap.add_argument("--export", metavar="FILE", help="write a self-contained HTML report")
    ap.add_argument("--test-connection", action="store_true", help="check the database connection")
    ap.add_argument("--demo", action="store_true", help="fill the cache with made-up data")
    ap.add_argument("--save-password", action="store_true",
                    help="store the SQL login password encrypted for your Windows account, then test the connection")
    ap.add_argument("--no-browser", action="store_true", help="don't open a browser when serving")
    ap.add_argument("--config", metavar="FILE", help="settings file (default config.ini), e.g. config_ieg.ini for another company")
    for key in PARAMS:
        ap.add_argument("--" + key.replace("_", "-"), type=int, dest=key)
    args = ap.parse_args(argv)
    global CONFIG_FILE
    if args.config:
        CONFIG_FILE = Path(args.config) if Path(args.config).is_absolute() else ROOT / args.config
        if not CONFIG_FILE.exists():
            ap.error(f"config file not found: {CONFIG_FILE}")
    cfg = load_config()
    apply_instance(cfg)
    if args.no_browser:
        cfg.set("server", "open_browser", "no")

    try:
        params = clean_params({k: getattr(args, k) for k in PARAMS}, default_params(cfg))
    except ValueError as ex:
        ap.error(str(ex))

    if args.save_password:
        import getpass
        user = cfg.get("database", "username", fallback="").strip()
        pw = getpass.getpass(f"SQL password for {user or 'the SQL login'} (typing is hidden): ")
        if not pw:
            print("No password entered. Nothing saved.")
            return 1
        if getpass.getpass("Type it again to confirm: ") != pw:
            print("The two entries didn't match. Nothing saved.")
            return 1
        if pw != pw.strip():
            print("Note: the password starts or ends with a space. If that's from pasting, run this again and type it.")
        save_password(pw)
        global _password_cache
        _password_cache = pw
        print(f"Saved, encrypted for your Windows account only, in {PASSWORD_FILE.relative_to(ROOT)}")
        args.test_connection = True

    if args.test_connection:
        import pyodbc
        print(f"Connecting to {source_label(cfg)} ...")
        try:
            with pyodbc.connect(connection_string(cfg), timeout=15) as conn:
                row = conn.cursor().execute("SELECT @@SERVERNAME, DB_NAME(), SUSER_SNAME()").fetchone()
            print(f"OK: server {row[0]}, database {row[1]}, login {row[2]}")
            return 0
        except Exception as ex:
            print("FAILED: " + friendly_error(ex))
            if "18456" not in str(ex) or not cfg.get("database", "database", fallback="").strip():
                return 1
        # Login failed: check whether the login works without naming a database.
        dbname = cfg.get("database", "database").strip()
        print(f"\nRetrying without database '{dbname}' to narrow it down ...")
        cfg.set("database", "database", "")
        try:
            with pyodbc.connect(connection_string(cfg), timeout=15) as conn:
                cur = conn.cursor()
                row = cur.execute("SELECT DB_NAME(), SUSER_SNAME()").fetchone()
                dbs = [r[0] for r in cur.execute(
                    "SELECT name FROM sys.databases WHERE HAS_DBACCESS(name) = 1 ORDER BY name").fetchall()]
            print(f"The login and password are fine (default database {row[0]}).")
            print(f"The problem is database '{dbname}': this login can't open it.")
            print("Databases this login can open: " + ", ".join(dbs))
            print("Put the right one on the 'database =' line in config.ini.")
        except Exception as ex2:
            print("Still failed, so the password (or the login itself) is the problem.")
            print("Run Save Password.bat again and type the password SSMS uses for mgambrell_read.")
            print("Detail: " + str(ex2)[:300])
        return 1

    if args.demo:
        t0 = time.time()
        sets = demo_sets(params)
        payload = build_payload(sets, params, "DEMO DATA (made up)", time.time() - t0)
        save_outputs(sets, payload, history=False)
        demo_history(payload)
        DEMO_MARKER.write_text("history.csv holds made-up data; the next real refresh replaces it\n")
        print("Demo data written. Run a real refresh to replace it.")

    if args.refresh:
        t0 = time.time()
        try:
            sets = run_query(cfg, params)
        except Exception as ex:
            print("FAILED: " + friendly_error(ex))
            return 1
        sets["_modules"] = run_modules(cfg, params)
        payload = build_payload(sets, params, source_label(cfg), time.time() - t0)
        save_outputs(sets, payload)
        print(f"Done in {time.time() - t0:,.0f}s")

    if args.export:
        payload = load_cache()
        if not payload:
            print("No cached data. Run with --refresh first.")
            return 1
        out = Path(args.export)
        out.write_text(render_static_html(payload), encoding="utf-8")
        print(f"Wrote {out.resolve()}")

    if not (args.refresh or args.export or args.demo):
        serve(cfg)
    return 0


if __name__ == "__main__":
    sys.exit(main())
